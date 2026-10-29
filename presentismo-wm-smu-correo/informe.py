"""Informe semanal "Act. Presentismo WM-SMU": arma el Excel y lo manda por correo.

  python informe.py                      manda a los destinatarios de destinatarios.json
  SOLO_A=a@x.cl,b@y.cl python informe.py prueba: manda solo a esos correos (sin CC) y
                                         antepone [Prueba] al asunto
  SOLO_ARCHIVO=ruta.xlsx python informe.py   guarda el Excel ahí y no envía nada
  VERIFICAR_HORA=10 python informe.py    sale sin hacer nada si en Chile no son las 10
                                         (el cron dispara dos veces por el horario de verano)

Semanas: las tres ISO que terminan en la semana del último día cerrado (ayer).
Hoy no entra en los % porque sus marcaciones están a medias.
"""
import os, re, sys, json, smtplib, ssl, urllib.request, urllib.parse, datetime as dt
from collections import defaultdict
from email.message import EmailMessage
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
CL = ZoneInfo("America/Santiago")
NOMBRE_ARCHIVO = "PRESENTISMO WM-SMU.xlsx"
REMITENTE = "supervisionecrapp@gmail.com"
URL = os.environ.get("SUPABASE_URL", "https://lbwwnrsbgaxjulpfbwdz.supabase.co") + "/rest/v1/"
COLS = ["horas_lunes", "horas_martes", "horas_miercoles", "horas_jueves", "horas_viernes", "horas_sabado", "horas_domingo"]
ZONAS = ["VIÑA", "TALCA", "IQUIQUE", "ARICA"]

AHORA = dt.datetime.now(CL)
HOY = AHORA.date()
if os.environ.get("VERIFICAR_HORA") and AHORA.hour != int(os.environ["VERIFICAR_HORA"]):
    print(f"En Chile son las {AHORA:%H:%M}, no las {os.environ['VERIFICAR_HORA']}: nada que hacer.")
    sys.exit(0)

def iso(d): return d.isocalendar()[:2]                 # (año ISO, semana)
def dias_sem(lunes): return [lunes + dt.timedelta(days=i) for i in range(7)]
REF = HOY - dt.timedelta(days=1)
LUN_C = REF - dt.timedelta(days=REF.weekday())
LUNES = [LUN_C - dt.timedelta(days=14), LUN_C - dt.timedelta(days=7), LUN_C]
WA, WB, WC = (dias_sem(l) for l in LUNES)
WCH = [d for d in WC if d < HOY]                        # semana en curso, solo días cerrados
SA, SB, SC = (iso(l)[1] for l in LUNES)
C_PARCIAL = len(WCH) < 7
D_DESDE = LUNES[0].isoformat()
ANIOS = sorted({iso(l)[0] for l in LUNES} | {iso(d)[0] for d in WC})
SEMANAS = sorted({iso(l)[1] for l in LUNES})

# ---------- Supabase
def _clave():
    k = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if k: return k
    p = os.path.join(HERE, "..", "presentismo-sync", ".env.local")      # corrida local
    for linea in open(p, encoding="utf-8"):
        m = re.match(r"SUPABASE_SERVICE_ROLE_KEY=(.*)", linea.strip())
        if m: return m.group(1).strip().strip('"').strip("'")
    raise SystemExit("Falta SUPABASE_SERVICE_ROLE_KEY")
KEY = _clave()

def get(tabla, params):
    params = dict(params)
    if "order" not in params: params["order"] = "sala_id,fecha" if tabla == "presentismo_horas_realizadas_por_sala_dia" else "id"
    elif tabla in ("presentismo_registros", "presentismo_smu_registros"): params["order"] += ",id"   # paginación estable
    out, off = [], 0
    while True:
        q = urllib.parse.urlencode(params, safe=",.()*:")
        req = urllib.request.Request(f"{URL}{tabla}?{q}", headers={"apikey": KEY, "Authorization": f"Bearer {KEY}", "Range-Unit": "items", "Range": f"{off}-{off+999}"})
        with urllib.request.urlopen(req, timeout=60) as r: filas = json.load(r)
        out += filas
        if len(filas) < 1000: return out
        off += 1000

def cargar():
    zonas = {z["id"]: z["nombre"] for z in get("zonas", {"select": "id,nombre"})}
    hold = {h["id"]: h["nombre"] for h in get("holdings", {"select": "id,nombre"})}
    cad = {c["id"]: c["nombre"] for c in get("cadenas", {"select": "id,nombre"})}
    prof = {p["id"]: p["full_name"] for p in get("profiles", {"select": "id,full_name"})}
    salas = get("salas", {"select": "id,sap,codigo_cadena,nombre_geo,holding_id,cadena_id,supervisor_id,coordinador_id,zona_id", "is_active": "eq.true"})
    for s in salas:
        s["zona"] = zonas.get(s["zona_id"]) or "-"; s["holding"] = hold.get(s["holding_id"]); s["cadena"] = cad.get(s["cadena_id"])
        s["sup"] = prof.get(s["supervisor_id"]) or "Sin supervisor"
    wm = [s for s in salas if s["holding"] == "WALMART" and s["cadena"] != "ACUENTA"]
    smu = [s for s in salas if s["holding"] == "SMU" and (s["cadena"] or "").strip().upper() not in ("S10", "SUPER 10")]
    ids = lambda l: ",".join(s["id"] for s in l)
    an, sm_ = "in.(" + ",".join(map(str, ANIOS)) + ")", "in.(" + ",".join(map(str, SEMANAS)) + ")"
    cols_obj = "sala_id,anio,semana," + ",".join(COLS)
    return {
        "wm": wm, "smu": smu,
        "pres_obj": get("presentismo_horas_objetivo", {"select": cols_obj, "sala_id": f"in.({ids(wm)})", "anio": an, "semana": sm_}),
        "pres_rea": get("presentismo_horas_realizadas_por_sala_dia", {"select": "sala_id,fecha,horas_realizadas", "sala_id": f"in.({ids(wm)})", "fecha": f"gte.{D_DESDE}"}),
        "pres_reg": get("presentismo_registros", {"select": "sala_id,nombre_local,rut_persona,nombre_persona,cargo,entrada,salida,horas", "sala_id": f"in.({ids(wm)})", "cargo": "eq.REPONEDOR", "entrada": f"gte.{D_DESDE}T00:00:00", "order": "entrada"}),
        "smu_cum": get("cumplimiento_smu", {"select": "sala_id,fecha,horas_cumplidas", "sala_id": f"in.({ids(smu)})", "fecha": f"gte.{D_DESDE}"}),
        "smu_obj": get("cumplimiento_smu_horas_objetivo", {"select": cols_obj, "sala_id": f"in.({ids(smu)})", "anio": an, "semana": sm_}),
        "smu_reg": get("presentismo_smu_registros", {"select": "sala_id,rut,nombre_persona,acceso,fecha", "sala_id": f"in.({ids(smu)})", "fecha": f"gte.{D_DESDE}", "order": "fecha"}),
        "vp": get("teamcore_venta_perdida", {"select": "sala_id,anio,semana,venta_clp_dia,venta_perdida_operacional", "sala_id": f"in.({ids(wm)})", "anio": an, "semana": sm_}),
    }

# ---------- estilos
NAVY = "1F2A44"
hdr_fill = PatternFill("solid", fgColor=NAVY); hdr_font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
sub_fill = PatternFill("solid", fgColor="E7EAF1"); bold = Font(bold=True, name="Calibri", size=10); base = Font(name="Calibri", size=10)
thin = Side(style="thin", color="C9CED8"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
RED = PatternFill("solid", fgColor="F8CFCF"); YEL = PatternFill("solid", fgColor="FBEFB5"); GRN = PatternFill("solid", fgColor="CFEBD3")
def cpf(p):   # colorForPct del portal: >=90 verde, >80 amarillo, resto rojo
    return None if p is None else GRN if p >= 0.9 else YEL if p > 0.8 else RED

def header(ws, row, labels, widths=None):
    for c, l in enumerate(labels, 1):
        x = ws.cell(row=row, column=c, value=l); x.fill = hdr_fill; x.font = hdr_font; x.border = box
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
def put(ws, row, col, val, fmt=None, fill=None, font=None, align=None):
    x = ws.cell(row=row, column=col, value=val); x.border = box; x.font = font or base
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    if align: x.alignment = Alignment(horizontal=align)
    return x
def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=14, name="Calibri", color=NAVY)
    if sub: ws["A2"] = sub; ws["A2"].font = Font(italic=True, size=9, name="Calibri", color="666666")

# ---------- presentismo / cumplimiento: suma sobre suma, real topado por día (como Difusión)
class Fuente:
    def __init__(self, salas, objetivo, realizado):
        self.salas = salas
        self.obj = {(o["sala_id"], o["anio"], o["semana"]): o for o in objetivo}
        self.rea = realizado
    def target(self, sid, d):
        a, s = iso(d); o = self.obj.get((sid, a, s)); v = o[COLS[d.weekday()]] if o else None
        return None if v is None else float(v)
    def real(self, sid, d): return float(self.rea.get((sid, d.isoformat()), 0) or 0)
    def pct(self, sids, dias):
        t = r = 0.0; hay = False
        for d in dias:
            for sid in sids:
                tg = self.target(sid, d)
                if tg is None: continue
                hay = True; t += tg; r += min(self.real(sid, d), tg)
        return r / t if hay and t > 0 else None

def armar(DATA):
    FW = Fuente(DATA["wm"], DATA["pres_obj"], {(r["sala_id"], r["fecha"]): r["horas_realizadas"] for r in DATA["pres_rea"]})
    smu_rea = {}
    for r in DATA["smu_cum"]: smu_rea[(r["sala_id"], r["fecha"])] = r["horas_cumplidas"]     # como el portal: gana la última fila
    FS = Fuente(DATA["smu"], DATA["smu_obj"], smu_rea)
    wb = Workbook()
    ws0 = wb.active; ws0.title = "General"

    # ===== Presentismo WM / Cumplimiento SMU: una hoja, todas las zonas hacia abajo
    def hoja_difusion(nombre_hoja, f, tit):
        ordz = {z: i for i, z in enumerate(ZONAS)}
        salas = sorted(f.salas, key=lambda s: (ordz.get(s["zona"], 9), s["sup"], s["nombre_geo"]))
        ws = wb.create_sheet(nombre_hoja)
        ws["A1"] = f"{tit} — todas las zonas"; ws["A1"].font = Font(bold=True, size=14, name="Calibri", color=NAVY)
        ws["A2"] = (f"Semanas {SA}, {SB} y {SC} (cerradas hasta el {HOY - dt.timedelta(days=1):%d/%m/%Y}; hoy no entra en los %) · "
                    "colores: ≥90% verde, 81–89% amarillo, ≤80% rojo")
        ws["A2"].font = Font(italic=True, size=9, name="Calibri", color="666666")
        NC = 19
        for c, l in enumerate(["Zona", "Sup", "SAP", "Cod.Cad", "Local"], 1):
            ws.merge_cells(start_row=4, start_column=c, end_row=5, end_column=c)
            x = ws.cell(row=4, column=c, value=l); x.fill = hdr_fill; x.font = hdr_font; x.alignment = Alignment(horizontal="center", vertical="center")
        grupos = ((6, ["L-S", "D", "Total"], SA), (9, ["L-S", "D", "Total"], SB), (12, [d.strftime("%d/%m") for d in WC] + ["Total"], SC))
        for c0, lbls, sem_ in grupos:
            ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + len(lbls) - 1)
            x = ws.cell(row=4, column=c0, value=sem_); x.fill = PatternFill("solid", fgColor="3A4A75"); x.font = hdr_font; x.alignment = Alignment(horizontal="center")
            for k, l in enumerate(lbls):
                y = ws.cell(row=5, column=c0 + k, value=l); y.fill = hdr_fill; y.font = hdr_font; y.alignment = Alignment(horizontal="center")
        for c in range(1, NC + 1):
            for r in (4, 5): ws.cell(row=r, column=c).border = box
        for c, w in enumerate([10, 18, 9, 8, 40] + [8] * 14, 1): ws.column_dimensions[get_column_letter(c)].width = w
        def fila(sids):
            return ([f.pct(sids, WA[:6]), f.pct(sids, WA[6:]), f.pct(sids, WA), f.pct(sids, WB[:6]), f.pct(sids, WB[6:]), f.pct(sids, WB)]
                    + [f.pct(sids, [d]) if d < HOY else None for d in WC] + [f.pct(sids, WCH)])
        TOT_COLS = (8, 11, 19)
        row = 6
        for s in salas:
            vals = [s["zona"].title(), s["sup"], s["sap"], s["codigo_cadena"], s["nombre_geo"]] + fila([s["id"]])
            for c, v in enumerate(vals, 1):
                put(ws, row, c, v if v is not None or c < 6 else "-", '0%' if c >= 6 else None, cpf(v) if c >= 6 else None,
                    bold if c in TOT_COLS else base, "center" if c not in (2, 5) else None)
            row += 1
        ws.auto_filter.ref = f"A5:{get_column_letter(NC)}{row-1}"
        row += 1                                                       # totales fuera del rango filtrado
        for nombre, zs in [(z.title(), [z]) for z in ZONAS] + [("TOTAL GENERAL", ZONAS)]:
            sids = [s["id"] for s in salas if s["zona"] in zs]
            if not sids: continue
            tot = nombre == "TOTAL GENERAL"
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            for c in range(1, 6):
                ws.cell(row=row, column=c).border = box; ws.cell(row=row, column=c).fill = sub_fill
            ws.cell(row=row, column=1, value=f"TOTALES {nombre}" if not tot else nombre).font = bold
            ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
            for c, v in enumerate(fila(sids), 6):
                put(ws, row, c, v if v is not None else "-", '0%', cpf(v), bold, "center")
            row += 1
        ws.freeze_panes = "F6"; ws.sheet_view.showGridLines = False
    hoja_difusion("Presentismo WM", FW, "Presentismo WM")
    hoja_difusion("Cumplimiento SMU", FS, "Cumplimiento SMU (Unimarc y Alvi)")

    # ===== General: resumen por zona
    vp = defaultdict(lambda: defaultdict(lambda: [0.0, 0.0]))        # sala -> (año, semana) -> [venta, perdida operacional]
    for r in DATA["vp"]:
        a = vp[r["sala_id"]][(r["anio"], r["semana"])]
        a[0] += float(r["venta_clp_dia"] or 0); a[1] += float(r["venta_perdida_operacional"] or 0)
    ws = ws0
    sufC = " (parcial)" if C_PARCIAL else ""
    title(ws, "Presentismo WM-SMU — resumen general por zona",
          f"Datos al {HOY:%d/%m/%Y}. Semana {SC} cerrada hasta el {HOY - dt.timedelta(days=1):%d/%m}. Presentismo/cumplimiento con tope diario; venta perdida = operacional ÷ venta (solo salas Walmart).")
    header(ws, 4, ["Zona", "Salas WM", f"Pres WM s{SA}", f"Pres WM s{SB}", f"Pres WM s{SC}", "Salas SMU", f"Cumpl SMU s{SA}", f"Cumpl SMU s{SB}", f"Cumpl SMU s{SC}",
                   f"Venta perdida $ s{SB}", f"% s{SA}", f"% s{SB}", f"% s{SC}{sufC}"], [14, 9, 11, 11, 11, 9, 12, 12, 12, 16, 9, 9, 12])
    def fila_general(zs):
        wm_ids = [s["id"] for s in DATA["wm"] if s["zona"] in zs]; sm_ids = [s["id"] for s in DATA["smu"] if s["zona"] in zs]
        def vv(w, k): return sum(vp[i][iso(LUNES[w])][k] for i in wm_ids)
        pr = lambda w: (vv(w, 1) / vv(w, 0)) if vv(w, 0) else None
        return [len(wm_ids), FW.pct(wm_ids, WA), FW.pct(wm_ids, WB), FW.pct(wm_ids, WCH), len(sm_ids), FS.pct(sm_ids, WA), FS.pct(sm_ids, WB), FS.pct(sm_ids, WCH),
                vv(1, 1), pr(0), pr(1), pr(2)]
    fmts = [None, '0%', '0%', '0%', None, '0%', '0%', '0%', '#,##0', '0.0%', '0.0%', '0.0%']
    resumen = {}
    for i, z in enumerate(ZONAS + ["TOTAL GENERAL"], 5):
        vals = fila_general([z] if z in ZONAS else ZONAS); resumen[z] = vals
        tot = z not in ZONAS
        put(ws, i, 1, z.title() if not tot else z, font=bold if tot else base, fill=sub_fill if tot else None)
        for c, (v, fm) in enumerate(zip(vals, fmts), 2):
            fill = cpf(v) if c in (3, 4, 5, 7, 8, 9) else (RED if v is not None and v >= 0.08 else YEL if v is not None and v >= 0.04 else None) if c in (11, 12, 13) else None
            put(ws, i, c, v if v is not None else "-", fm, fill, bold if tot else base, "center")
    ws.freeze_panes = "B5"; ws.auto_filter.ref = "A4:M9"

    # ===== detalle por día y marcaciones
    def tabla(nombre, encabezado, filas, anchos, fmts=None, fills=None):
        w = wb.create_sheet(nombre[:31])
        for c, l in enumerate(encabezado, 1):
            x = w.cell(row=1, column=c, value=l); x.fill = hdr_fill; x.font = hdr_font; x.border = box
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for c, a in enumerate(anchos, 1): w.column_dimensions[get_column_letter(c)].width = a
        for r, fl in enumerate(filas, 2):
            for c, v in enumerate(fl, 1):
                x = w.cell(row=r, column=c, value=v); x.font = base; x.border = box
                if fmts and fmts.get(c): x.number_format = fmts[c]
                if fills:
                    fi = fills(c, v)
                    if fi: x.fill = fi
        w.freeze_panes = "A2"; w.auto_filter.ref = w.dimensions
    DIAS = WA + WB + WCH
    def x_dia(f, salas, nombre):
        filas = [[s["nombre_geo"], s["sup"], s["zona"]] + [f.pct([s["id"]], [d]) for d in DIAS]
                 for s in sorted(salas, key=lambda s: (s["zona"], s["sup"], s["nombre_geo"]))]
        tabla(nombre, ["Sala", "Supervisor", "Zona"] + [d.strftime("%d/%m") for d in DIAS], filas, [40, 18, 18] + [7] * len(DIAS),
              {c: '0%' for c in range(4, 4 + len(DIAS))}, lambda c, v: cpf(v) if c >= 4 else None)
    def detalle(f, salas, nombre, etiqueta):
        filas = []
        for s in sorted(salas, key=lambda s: (s["zona"], s["sup"], s["nombre_geo"])):
            for d in DIAS:
                tg = f.target(s["id"], d)
                if tg is None: continue
                rl = f.real(s["id"], d)
                filas.append([s["nombre_geo"] if f"({s['sap']})" in s["nombre_geo"] else f"{s['nombre_geo']} ({s['sap']})", s["sup"], d.strftime("%d/%m/%Y"),
                              tg / 24, rl / 24, min(rl, tg) / tg if tg > 0 else None])
        tabla(nombre, ["Sala", "Supervisor", "Fecha", "Horas Objetivo", etiqueta, "%"], filas, [44, 18, 12, 14, 16, 8],
              {4: '[h]:mm', 5: '[h]:mm', 6: '0%'}, lambda c, v: cpf(v) if c == 6 else None)
    x_dia(FW, DATA["wm"], "Pres WM x dia")
    detalle(FW, DATA["wm"], "Pres WM detalle", "Horas Realizadas")

    sid2 = {s["id"]: s for s in DATA["wm"]}
    def hm(ts):
        return dt.datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(CL).replace(tzinfo=None) if ts else None
    filas = []
    for r in sorted(DATA["pres_reg"], key=lambda r: ((r["nombre_local"] or sid2[r["sala_id"]]["nombre_geo"]), r["entrada"] or "")):
        s = sid2[r["sala_id"]]; e = hm(r["entrada"]); sl = hm(r["salida"])
        if e is None or e.date() > HOY: continue
        # Fecha y Semana salen de la entrada
        filas.append([r["nombre_local"] or s["nombre_geo"], s["sap"], s["sup"], r["rut_persona"], r["nombre_persona"], e.date(), iso(e.date())[1], e,
                      sl if sl else "Sin marcar", r["horas"] / 24 if r["horas"] is not None else None])
    tabla("Marcaciones WM", ["Sala", "SAP", "Supervisor", "RUT", "Nombre", "Fecha", "Semana", "Entrada", "Salida", "Horas"], filas, [40, 9, 18, 13, 32, 11, 9, 17, 17, 8],
          {6: 'dd/mm/yyyy', 8: 'dd/mm/yyyy hh:mm', 9: 'dd/mm/yyyy hh:mm', 10: '[h]:mm'})
    n_wm = len(filas)

    x_dia(FS, DATA["smu"], "Cumpl SMU x dia")
    detalle(FS, DATA["smu"], "Cumpl SMU detalle", "Horas Cumplidas")
    sm = {s["id"]: s for s in DATA["smu"]}
    filas = []
    for r in sorted(DATA["smu_reg"], key=lambda r: (sm[r["sala_id"]]["nombre_geo"], r["fecha"], r["rut"])):
        if r["fecha"] > HOY.isoformat(): continue
        s = sm[r["sala_id"]]
        filas.append([s["nombre_geo"], s["sap"], s["sup"], r["rut"], r["nombre_persona"] or "-", r["fecha"], r["acceso"]])
    tabla("Marcas SMU", ["Sala", "SAP", "Supervisor", "RUT", "Nombre", "Fecha", "Acceso"], filas, [40, 9, 18, 13, 32, 12, 10])

    orden = ["General", "Presentismo WM", "Cumplimiento SMU"]
    orden += [n for n in wb.sheetnames if n not in orden]
    wb._sheets = [wb[n] for n in orden]
    wb.properties.title = "PRESENTISMO WM-SMU"
    return wb, resumen, n_wm

# ---------- correo
def _p(v): return "-" if v is None else f"{v*100:.1f}%".replace(".", ",")
def cuerpo(resumen):
    celda = lambda v: f'<td style="padding:4px 12px;border-bottom:1px solid #eef0f2;text-align:center;font-weight:600;color:{"#6b7280" if v is None else "#15803d" if v >= .9 else "#b45309" if v > .8 else "#b91c1c"}">{_p(v)}</td>'
    th = 'style="padding:4px 12px;border-bottom:2px solid #d1d5db;font-size:12px;color:#4b5563"'
    filas = ""
    for z in ZONAS + ["TOTAL GENERAL"]:
        v = resumen[z]; t = z == "TOTAL GENERAL"
        filas += (f'<tr style="{"background:#f3f4f6;font-weight:700" if t else ""}"><td style="padding:4px 12px;border-bottom:1px solid #eef0f2">{z.title() if not t else z}</td>'
                  + "".join(celda(v[i]) for i in (1, 2, 3)) + "".join(celda(v[i]) for i in (5, 6, 7)) + "</tr>")
    return f"""<div style="font:14px/1.5 Segoe UI,Arial,sans-serif;color:#1f2937">
<p>Hola,</p>
<p>Adjunto la actualización de presentismo de Walmart y cumplimiento de horas de Unimarc/Alvi (SMU), con las semanas {SA}, {SB} y {SC}
(cerradas hasta el {HOY - dt.timedelta(days=1):%d/%m/%Y}). El archivo trae el detalle por sala, por día y las marcaciones.</p>
<table style="border-collapse:collapse;font-size:13px">
<tr><th {th}></th><th {th} colspan="3">Presentismo WM</th><th {th} colspan="3">Cumplimiento SMU</th></tr>
<tr><th {th}>Zona</th>{"".join(f'<th {th}>S{s}</th>' for s in (SA, SB, SC)) * 2}</tr>
{filas}</table>
<p style="color:#6b7280;font-size:12px">Verde ≥90%, amarillo 81–89%, rojo ≤80%. Se calcula con tope diario (lo que sobra de horas un día no compensa otro).
{"La semana " + str(SC) + " aún está en curso." if C_PARCIAL else ""}</p>
<p>Saludos,<br>Supervisión ECR</p></div>"""

def enviar(ruta_xlsx, resumen):
    cfg = json.load(open(os.path.join(HERE, "destinatarios.json"), encoding="utf-8"))
    solo = [x.strip() for x in os.environ.get("SOLO_A", "").split(",") if x.strip()]
    para, cc = (solo, []) if solo else (cfg["para"], cfg["cc"])
    pw = re.sub(r"\s+", "", os.environ.get("GMAIL_APP_PASSWORD", ""))
    if not pw: raise SystemExit("Falta GMAIL_APP_PASSWORD")
    m = EmailMessage()
    m["From"] = f"Informes Supervisión ECR <{REMITENTE}>"
    m["To"] = ", ".join(para)
    if cc: m["Cc"] = ", ".join(cc)
    m["Subject"] = ("[Prueba] " if solo else "") + "Act. Presentismo WM-SMU"
    m.set_content("Adjunto la actualización de Presentismo WM-SMU.")
    m.add_alternative(cuerpo(resumen), subtype="html")
    m.add_attachment(open(ruta_xlsx, "rb").read(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=NOMBRE_ARCHIVO)
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=ssl.create_default_context()) as s:
        s.login(REMITENTE, pw)
        s.send_message(m)
    print(f"Enviado a {len(para)} destinatarios" + (f" + {len(cc)} en copia" if cc else "") + f": {', '.join(para + cc)}")

if __name__ == "__main__":
    print(f"Semanas {SA}/{SB}/{SC}; hoy {HOY}; días cerrados de la semana {SC}: {len(WCH)}")
    wb, resumen, n_wm = armar(cargar())
    destino = os.environ.get("SOLO_ARCHIVO") or os.path.join(os.environ.get("RUNNER_TEMP") or os.environ.get("TEMP") or "/tmp", NOMBRE_ARCHIVO)
    wb.save(destino)
    print(f"Excel: {destino} ({n_wm} marcaciones WM)")
    if not os.environ.get("SOLO_ARCHIVO"): enviar(destino, resumen)
